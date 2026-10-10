"""Excel (.xlsx) -> Markdown 转换器。"""
from __future__ import annotations

import os

from .base import ConversionError, Converter


class XlsxConverter(Converter):
    extensions = (".xlsx",)

    def convert(self, path: str) -> str:
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise ConversionError("缺少 openpyxl 库") from e

        try:
            wb = load_workbook(path, data_only=True, read_only=True)
        except Exception as e:
            raise ConversionError(f"无法打开 xlsx 文件: {e}") from e

        title = os.path.splitext(os.path.basename(path))[0]
        lines: list[str] = [f"# {title}", ""]

        for ws in wb.worksheets:
            lines.append(f"## 工作表：{ws.title}")
            lines.append("")
            md_table = self._sheet_to_md(ws)
            if md_table:
                lines.append(md_table)
                lines.append("")

        wb.close()
        return "\n".join(lines).strip() + "\n"

    def _sheet_to_md(self, ws) -> str:
        rows_iter = ws.iter_rows(values_only=True)
        rows: list[list[str]] = []
        col_count = 0
        for raw in rows_iter:
            if not any(v not in (None, "") for v in raw):
                # 跳过全空行
                continue
            cells = [self._cell_to_str(v) for v in raw]
            col_count = max(col_count, len(cells))
            rows.append(cells)
        if not rows:
            return ""

        # 补齐每行列数，使表格对齐
        for r in rows:
            while len(r) < col_count:
                r.append("")

        out = []
        for i, r in enumerate(rows):
            out.append("| " + " | ".join(r) + " |")
            if i == 0:
                out.append("|" + "|".join(["---"] * col_count) + "|")
        return "\n".join(out)

    def _cell_to_str(self, v) -> str:
        if v is None:
            return ""
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return self._md_escape(self._clean(str(v)))
