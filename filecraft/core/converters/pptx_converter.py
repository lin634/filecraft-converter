"""PowerPoint (.pptx) -> Markdown 转换器。"""
from __future__ import annotations

import os

from .base import ConversionError, Converter


class PptxConverter(Converter):
    extensions = (".pptx",)

    def convert(self, path: str) -> str:
        try:
            from pptx import Presentation
        except ImportError as e:
            raise ConversionError("缺少 python-pptx 库") from e

        try:
            prs = Presentation(path)
        except Exception as e:
            raise ConversionError(f"无法打开 pptx 文件: {e}") from e

        title = os.path.splitext(os.path.basename(path))[0]
        lines: list[str] = [f"# {title}", ""]

        for idx, slide in enumerate(prs.slides, start=1):
            lines.append(f"## 幻灯片 {idx}")
            lines.append("")
            for shape in slide.shapes:
                md = self._shape_to_md(shape)
                if md:
                    lines.append(md)
                    lines.append("")
                if shape.has_table:
                    lines.append(self._table_to_md(shape.table))
                    lines.append("")
            # 备注
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = self._clean(slide.notes_slide.notes_text_frame.text)
                if notes:
                    lines.append(f"> 备注：{notes}")
                    lines.append("")
            lines.append("---")
            lines.append("")

        return "\n".join(lines).strip() + "\n"

    # ---- helpers --------------------------------------------------------
    def _shape_to_md(self, shape) -> str:
        out: list[str] = []
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = self._clean(para.text)
                if text:
                    level = para.level or 0
                    prefix = "  " * level + "- "
                    out.append(f"{prefix}{text}")
        if shape.shape_type == 13:  # Picture
            out.append(f"![图片]({shape.name})")
        return "\n".join(out)

    def _table_to_md(self, table) -> str:
        rows = []
        for i, row in enumerate(table.rows):
            cells = [self._md_escape(self._clean(c.text)) for c in row.cells]
            rows.append("| " + " | ".join(cells) + " |")
            if i == 0:
                rows.append("|" + "|".join(["---"] * len(cells)) + "|")
        return "\n".join(rows)
