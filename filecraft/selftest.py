"""冻结包自检。

用 PyInstaller 打包后，代码里函数内部的延迟导入（如 `from docx import Document`）
有可能会被漏掉。GUI 程序又看不到控制台输出，所以提供一个自检入口：
在打包后的 exe 上运行 `FileCraft.exe --selftest 输出文件.txt`，
它会真实走一遍四种格式的转换和 PDF 导出，把结果写进文件。

用法：
    python selftest.py [输出文件]      # 源码模式
    FileCraft.exe --selftest out.txt   # 打包后
"""
from __future__ import annotations

import io
import os
import sys
import tempfile
import traceback


def _run_checks() -> list[tuple[str, bool, str]]:
    results: list[tuple[str, bool, str]] = []

    def check(name: str, fn):
        try:
            detail = fn() or ""
            results.append((name, True, str(detail)))
        except Exception as e:  # noqa: BLE001
            results.append((name, False, f"{type(e).__name__}: {e}\n{traceback.format_exc()}"))

    tmp = tempfile.mkdtemp(prefix="filecraft_selftest_")

    # ---- 依赖导入 ------------------------------------------------------
    def _imports():
        mods = ["PyQt6.QtWidgets", "docx", "pptx", "openpyxl", "pdfplumber",
                "markdown", "openai", "reportlab.pdfbase.ttfonts"]
        import importlib
        for m in mods:
            importlib.import_module(m)
        return f"{len(mods)} 个依赖模块全部导入成功"

    check("依赖模块导入", _imports)

    # ---- 造测试文档 ----------------------------------------------------
    docx_path = os.path.join(tmp, "t.docx")
    pptx_path = os.path.join(tmp, "t.pptx")
    xlsx_path = os.path.join(tmp, "t.xlsx")
    pdf_path = os.path.join(tmp, "t.pdf")

    def _make_docx():
        from docx import Document
        d = Document()
        d.add_heading("经营分析", level=1)
        d.add_paragraph("营业收入达到一亿元。")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "项目"
        t.cell(0, 1).text = "金额"
        t.cell(1, 0).text = "营收"
        t.cell(1, 1).text = "1亿"
        d.save(docx_path)
        return "已生成"

    check("生成 docx 测试文件", _make_docx)

    def _make_pptx():
        from pptx import Presentation
        p = Presentation()
        s = p.slides.add_slide(p.slide_layouts[0])
        s.shapes.title.text = "产品路线图"
        s.placeholders[1].text = "2026 规划"
        p.save(pptx_path)
        return "已生成"

    check("生成 pptx 测试文件", _make_pptx)

    def _make_xlsx():
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["区域", "销售额"])
        ws.append(["华东", 5200])
        wb.save(xlsx_path)
        return "已生成"

    check("生成 xlsx 测试文件", _make_xlsx)

    def _make_pdf():
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas
        from core.md_to_pdf import pick_cjk_font
        fp = pick_cjk_font()
        if fp:
            pdfmetrics.registerFont(TTFont("STTest", fp))
            font = "STTest"
        else:
            font = "Helvetica"
        c = canvas.Canvas(pdf_path, pagesize=A4)
        c.setFont(font, 14)
        c.drawString(70, 760, "项目验收通知")
        c.setFont(font, 11)
        c.drawString(70, 730, "验收时间为十月十五日。")
        c.save()
        return f"已生成（字体={os.path.basename(fp) if fp else 'Helvetica'}）"

    check("生成 pdf 测试文件（需中文字体）", _make_pdf)

    # ---- 四种格式转换 --------------------------------------------------
    def _mk_convert(path, expect, label):
        def _inner():
            from core.registry import convert_file
            md = convert_file(path)
            if expect not in md:
                raise AssertionError(f"转换结果中缺少预期内容 {expect!r}，实际: {md[:200]!r}")
            return f"{len(md)} 字符，含 {expect!r}"
        check(label, _inner)

    _mk_convert(docx_path, "经营分析", "转换 docx")
    _mk_convert(pptx_path, "产品路线图", "转换 pptx")
    _mk_convert(xlsx_path, "华东", "转换 xlsx")
    _mk_convert(pdf_path, "验收时间", "转换 pdf")

    # ---- 不支持的格式应被拒绝 ------------------------------------------
    def _reject():
        from core.registry import convert_file, is_supported
        if is_supported("x.doc"):
            raise AssertionError(".doc 不应被判定为支持")
        try:
            convert_file("x.doc")
        except Exception:
            return "旧格式已正确拒绝"
        raise AssertionError("convert_file 未拒绝 .doc")
    check("旧格式拒绝逻辑", _reject)

    # ---- Markdown -> PDF（含中文字体与嵌套列表）------------------------
    def _pdf_export():
        from core.md_to_pdf import markdown_to_pdf
        md = (
            "# 文档总结\n\n## 总体概述\n营业收入 **一亿元**，同比 +18%。\n\n"
            "## 要点\n\n- 营收增长\n  - 华东领先\n  - 华北持平\n\n"
            "| 区域 | 销售额 |\n|---|---|\n| 华东 | 5200 |\n"
        )
        out = os.path.join(tmp, "summary.pdf")
        markdown_to_pdf(md, out, title="文档总结")
        size = os.path.getsize(out)
        if size < 1000:
            raise AssertionError(f"PDF 过小: {size} 字节")
        # 反解验证中文确实写入
        import pdfplumber
        with pdfplumber.open(out) as pdf:
            text = "\n".join((p.extract_text() or "") for p in pdf.pages)
        for kw in ["文档总结", "总体概述", "营业收入", "华东", "一亿元"]:
            if kw not in text:
                raise AssertionError(f"PDF 中缺少 {kw!r}")
        return f"{size} 字节，中文反解校验通过"

    check("Markdown 导出 PDF（中文+嵌套列表）", _pdf_export)

    # ---- 长文本分块 ----------------------------------------------------
    def _chunking():
        from config import LLMConfig
        from core.summarizer import FileInfo, Summarizer
        cfg = LLMConfig(max_chars_per_chunk=100)
        s = Summarizer(cfg)
        chunks = s._split_chunks([FileInfo("a", "x" * 250), FileInfo("b", "y" * 50)])
        if len(chunks) < 3:
            raise AssertionError(f"分块数异常: {len(chunks)}")
        return f"切成 {len(chunks)} 块"
    check("长文本分块逻辑", _chunking)

    # ---- 配置读写 ------------------------------------------------------
    def _config():
        from config import CONFIG_DIR, load_config
        cfg = load_config()
        return f"配置目录 {CONFIG_DIR}，当前模型 {cfg.llm.model}"
    check("配置读取", _config)

    return results


def run_selftest(out_path: str | None = None) -> int:
    results = _run_checks()
    from version import APP_VERSION

    lines = [f"FileCraft 自检报告 v{APP_VERSION}", "=" * 52, ""]
    passed = failed = 0
    for name, ok, detail in results:
        if ok:
            passed += 1
            lines.append(f"[通过] {name}")
        else:
            failed += 1
            lines.append(f"[失败] {name}")
        for ln in str(detail).splitlines():
            lines.append(f"        {ln}")
        lines.append("")
    lines.append("=" * 52)
    lines.append(f"合计 {len(results)} 项：通过 {passed}，失败 {failed}")
    lines.append(f"运行环境: Python {sys.version.split()[0]}")
    lines.append(f"是否冻结包: {getattr(sys, 'frozen', False)}")
    if getattr(sys, "frozen", False):
        lines.append(f"可执行文件: {sys.executable}")
    report = "\n".join(lines)

    if out_path:
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(report)
    # 冻结的窗口版没有控制台，尝试打印也无妨
    try:
        print(report)
    except Exception:
        pass
    return 1 if failed else 0


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else None
    raise SystemExit(run_selftest(target))